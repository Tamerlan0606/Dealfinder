import io, os, re, tempfile
from urllib.parse import urljoin
import httpx

try:
    from pypdf import PdfReader
except Exception:
    PdfReader = None

try:
    from openpyxl import load_workbook
except Exception:
    load_workbook = None

DOC_RE = re.compile(r'href=["\']([^"\']+\.(?:pdf|xlsx?|xlsm)(?:\?[^"\']*)?)["\']', re.I)

def extract_document_links(base_url, html):
    links = []
    for raw in DOC_RE.findall(html or ""):
        u = urljoin(base_url, raw)
        if u not in links:
            links.append(u)
    return links[:12]

def _pdf_text(data):
    if not PdfReader:
        return ""
    try:
        reader = PdfReader(io.BytesIO(data))
        return "\n".join((p.extract_text() or "") for p in reader.pages[:80])
    except Exception:
        return ""

def _xlsx_text(data):
    if not load_workbook:
        return ""
    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        chunks = []
        for ws in wb.worksheets[:20]:
            for row in ws.iter_rows(values_only=True):
                vals = [str(v) for v in row if v is not None]
                if vals:
                    chunks.append(" | ".join(vals))
                if len(chunks) >= 30000:
                    break
            if len(chunks) >= 30000:
                break
        return "\n".join(chunks)
    except Exception:
        return ""

def fetch_and_extract(client, url):
    try:
        r = client.get(url, timeout=httpx.Timeout(12.0, connect=5.0))
        r.raise_for_status()
        ct = (r.headers.get("content-type") or "").lower()
        data = r.content
        if ".pdf" in url.lower() or "pdf" in ct:
            text = _pdf_text(data)
            kind = "pdf"
        elif re.search(r'\.(xlsx?|xlsm)(?:\?|$)', url, re.I) or "spreadsheet" in ct or "excel" in ct:
            text = _xlsx_text(data)
            kind = "xlsx"
        else:
            text = ""
            kind = "other"
        return {"url": url, "kind": kind, "bytes": len(data), "text": text[:500000]}
    except Exception as e:
        return {"url": url, "kind": "error", "error": str(e), "text": ""}

def collect_documents(url, timeout=20):
    if not url:
        return {"links": [], "documents": [], "combined_text": ""}
    try:
        with httpx.Client(timeout=httpx.Timeout(timeout, connect=8.0), follow_redirects=True,
                           headers={"User-Agent": "DealFinder/6.0"}) as client:
            r = client.get(url)
            r.raise_for_status()
            links = extract_document_links(str(r.url), r.text)
            docs = [fetch_and_extract(client, x) for x in links]
            text = "\n\n".join(d.get("text","") for d in docs if d.get("text"))
            return {"links": links, "documents": docs, "combined_text": text[:1000000]}
    except Exception as e:
        return {"links": [], "documents": [], "combined_text": "", "error": str(e)}
